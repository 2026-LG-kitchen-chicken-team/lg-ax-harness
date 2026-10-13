"""Reader rendering contracts use new synthetic fixtures, never business data."""
import copy
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "product-tools"))
from product_tools.contracts import ToolError
from product_tools import reporting


def example_plan():
    return {
        "title": "합성 자료로 확인한 보고서 생성",
        "summary": ["표와 차트는 계약 시험용 합성 값이다.", "고객 원문과 작성자의 설명을 구분한다."],
        "sources": [
            {"id": "internal-review-001", "kind": "review", "title": "시험 후기 원문",
             "text": '=1+1\n“문을 열 때 손잡이가 불편했다.”\n' + "원문 보존 확인. " * 95,
             "url": "https://example.com/review", "location": "후기 1 문단",
             "model": "모델 시험 A / 모델 버전 1 / 문서 버전 2 / 자료 버전 3", "synthetic": True},
            {"id": "internal-survey-002", "kind": "survey", "title": "시험 설문 원문",
             "text": "시험 응답 열 명. 항목 A 네 명, 항목 B 여섯 명.", "synthetic": True}
        ],
        "sections": [
            {"title": "응답 구성", "blocks": [
                {"type": "paragraph", "text": "시험 응답은 항목 A 네 명과 항목 B 여섯 명이다.",
                 "source_ids": ["internal-survey-002"]},
                {"type": "table", "title": "시험 응답과 비율", "columns": ["항목", "응답 수", "비율"],
                 "rows": [["항목 A", 4, "40%"], ["항목 B", 6, "60%"], ["합계", 10, "100%"]],
                 "total_rows": [2], "highlight_cells": [[1, 1]],
                 "source_ids": ["internal-survey-002"]},
                {"type": "chart", "title": "항목별 시험 응답", "chart_type": "bar",
                 "labels": ["항목 A", "항목 B"], "series": [{"name": "응답", "values": [4, 6]}],
                 "unit": "명", "caption": "시험 응답 열 명 중 항목 B는 여섯 명이다.",
                 "highlights": [[0, 1]], "source_ids": ["internal-survey-002"]},
                {"type": "chart", "title": "시험 응답 구성", "chart_type": "donut",
                 "labels": ["항목 A", "항목 B"], "series": [{"name": "응답", "values": [4, 6]}],
                 "unit": "명", "caption": "상호 배타적인 두 항목이 전체 열 명을 구성한다.",
                 "highlights": [[0, 1]], "source_ids": ["internal-survey-002"]}
            ]},
            {"title": "시간과 측정값", "blocks": [
                {"type": "chart", "title": "순서별 시험 측정값", "chart_type": "line",
                 "labels": ["첫째", "둘째", "셋째"], "series": [{"name": "측정값", "values": [-1, 2, 3]}],
                 "unit": "시험 단위", "caption": "음수와 양수도 같은 축에서 확인한다.",
                 "highlights": [[0, 2]], "source_ids": ["internal-survey-002"]},
                {"type": "chart", "title": "좌표별 시험 측정값", "chart_type": "scatter",
                 "labels": ["점 A", "점 B", "점 C"],
                 "series": [{"name": "시험 좌표", "x_values": [1, 3, 7], "values": [2, -1, 4]}],
                 "unit": "Y 시험 단위", "caption": "X 좌표 간격을 그대로 사용한다.",
                 "highlights": [[0, 2]], "source_ids": ["internal-survey-002"]},
                {"type": "paragraph", "text": "“문을 열 때 손잡이가 불편했다.”",
                 "source_ids": ["internal-review-001"]}
            ]}
        ],
        "unfulfilled_requests": [{"request": "실제 이미지 요청", "reason": "시험 입력에는 실제 이미지가 제공되지 않았다."}]
    }


def multi_source_plan():
    sources = [{"id": "internal-citation-"+str(i), "kind": "other",
                "title": "=SUM(A1:A2)" if i == 0 else "시험 근거 "+str(i+1),
                "text": "원문 기록 "+str(i+1), "synthetic": True} for i in range(5)]
    line = {"type":"chart","title":"두 계열의 합성 측정값","chart_type":"line",
            "labels":["첫째","둘째","셋째"],"unit":"시험 단위",
            "series":[{"name":"계열 A","values":[1,5,2]},{"name":"계열 B","values":[5,1,4]}],
            "source_ids":[s["id"] for s in sources],"highlights":[[1,2]],
            "caption":"합성 계열 A와 B가 둘째 구간에서 교차한다.\n회색 원과 실선은 A, 진한 회색 사각형과 점선은 B를 나타낸다."}
    scatter = copy.deepcopy(line)
    scatter.update(title="두 계열의 합성 좌표",chart_type="scatter",
                   caption="합성 좌표의 계열을 범례의 원과 사각형으로 구분한다.")
    scatter["series"][0]["x_values"] = [1,2,4]
    scatter["series"][1]["x_values"] = [1,3,4]
    return {"title":"출처와 복수 계열 계약 시험","summary":["계약 시험용 합성 자료다."],
            "sources":sources,"sections":[{"title":"교차하는 두 계열","blocks":[line,scatter]}]}


class ReportingValidationTests(unittest.TestCase):
    def test_multiple_series_have_distinct_matching_keys_and_red_highlight(self):
        from reportlab.graphics.shapes import Circle, Rect, Line, String
        from reportlab.lib.colors import HexColor
        for block in reporting._adapt_plan(multi_source_plan())["sections"][0]["blocks"]:
            d = reporting._chart_drawings(block,reporting._font())[0]
            points = [v for v in d.contents if isinstance(v,(Circle,Rect))
                      and (getattr(v,"r",None) == 3.5 or getattr(v,"width",None) == 7)]
            circles = [v for v in points if isinstance(v,Circle)]
            squares = [v for v in points if isinstance(v,Rect)]
            self.assertEqual(len(circles),4)  # three plot points plus the key
            self.assertEqual(len(squares),4)
            self.assertEqual(circles[0].fillColor,circles[-1].fillColor)
            self.assertEqual(squares[0].fillColor,squares[-1].fillColor)
            self.assertNotEqual(circles[-1].fillColor,squares[-1].fillColor)
            self.assertEqual(squares[2].fillColor,HexColor("#B42318"))
            self.assertEqual({v.text for v in d.contents if isinstance(v,String)} & {"계열 A","계열 B"},
                             {"계열 A","계열 B"})
            if block["kind"] == "line":
                strokes = [v for v in d.contents if isinstance(v,Line)
                           and v.strokeColor in (circles[-1].fillColor,squares[-1].fillColor)
                           and v.y1 > 65 and v.y2 > 65]
                self.assertTrue(any(v.strokeDashArray for v in strokes))
                self.assertTrue(any(not v.strokeDashArray for v in strokes))

    def test_copy_and_exact_schema(self):
        plan = example_plan()
        normalized = reporting.validate_report_plan(plan)
        self.assertEqual(plan, normalized)
        self.assertIsNot(plan, normalized)
        for bad in ({"extra": 1}, {"summary": "text"}):
            value = {**plan, **bad}
            with self.assertRaises(ToolError):
                reporting.validate_report_plan(value)

    def test_invalid_source_and_table_dimensions(self):
        for mutate in (
            lambda p: p["sources"].append(copy.deepcopy(p["sources"][0])),
            lambda p: p["sections"][0]["blocks"][0].update(source_ids=["unknown"]),
            lambda p: p["sections"][0]["blocks"][1]["rows"].append(["short"]),
            lambda p: p["sections"][0]["blocks"][1].update(total_rows=[99]),
            lambda p: p["sections"][0]["blocks"][1].update(highlight_cells=[[0, 9]]),
            lambda p: p["sections"][0]["blocks"][1]["rows"][0].__setitem__(1, True),
            lambda p: p["sources"][0].update(url="file:///secret"),
            lambda p: p["sources"][0].update(synthetic=None),
        ):
            value = example_plan()
            mutate(value)
            with self.assertRaises(ToolError):
                reporting.validate_report_plan(value)

    def test_finite_chart_source_and_dimensions(self):
        for mutate in (
            lambda b: b.update(source_ids=[]),
            lambda b: b["series"][0].update(values=[float("nan"), 1]),
            lambda b: b["series"][0].update(values=[float("inf"), 1]),
            lambda b: b["series"][0].update(values=[10**400, 1]),
            lambda b: b["series"][0].update(values=[1]),
            lambda b: b.update(highlights=[[99, 0]]),
        ):
            value = example_plan()
            mutate(value["sections"][0]["blocks"][2])
            with self.assertRaises(ToolError):
                reporting.validate_report_plan(value)
        value = example_plan()
        del value["sections"][1]["blocks"][1]["series"][0]["x_values"]
        with self.assertRaises(ToolError):
            reporting.validate_report_plan(value)
        value = example_plan()
        value["sections"][0]["blocks"][3]["series"][0]["values"] = [0, 0]
        with self.assertRaises(ToolError):
            reporting.validate_report_plan(value)

    def test_image_blocks_are_unsupported_for_every_path(self):
        with tempfile.TemporaryDirectory() as folder:
            bad = Path(folder)/"corrupt.png"
            bad.write_bytes(b"not an image")
            from PIL import Image
            valid = Path(folder)/"valid.png"
            Image.new("RGB", (20, 20), "white").save(valid)
            for path in (str(valid), "relative.png", str(bad), str(Path(folder)/".."/"escape.png"),
                         "https://example.com/image.png", "\\\\server\\share\\image.png"):
                value = example_plan()
                value["sections"][0]["blocks"].append({"type": "image", "path": path, "caption": "시험"})
                with self.assertRaises(ToolError):
                    reporting.validate_report_plan(value)


class ReportingRenderTests(unittest.TestCase):
    def test_excel_print_width_keeps_full_heading_and_no_blank_footer_row(self):
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
        plan = example_plan()
        blocks = copy.deepcopy(plan["sections"][0]["blocks"][1:3])
        blocks[1]["title"] = "대한민국 일반가구 중 1인가구 비중 · 2020~2025년 / 긴 제목 보존 검증"
        plan["sections"] = [{"title": "인쇄 검증", "blocks": blocks}]
        with tempfile.TemporaryDirectory() as temporary:
            result = reporting.render_report(plan, Path(temporary)/"reader")
            self.assertTrue(result["ready"], result)
            book = load_workbook(Path(result["directory"])/"01_자료집"/"보고서 도표.xlsx")
            table, chart, bibliography = book.worksheets
            last = max(cell.row for row in table for cell in row if cell.value is not None)
            self.assertEqual(range_boundaries(str(table.print_area).split("!")[-1])[3], len(blocks[0]["rows"])+6)
            self.assertGreater(last, len(blocks[0]["rows"])+6)
            self.assertEqual(bibliography.title, "도표 출처")
            self.assertTrue(bibliography["D7"].hyperlink)
            self.assertEqual(chart["B2"].value, blocks[1]["title"])
            self.assertTrue(chart["B2"].alignment.wrap_text)
            self.assertIn("B2:M2", [str(r) for r in chart.merged_cells.ranges])
            self.assertGreaterEqual(chart.row_dimensions[2].height, 32)

    def test_emphasis_quote_styles_and_literal_raw_sources(self):
        from docx import Document
        from openpyxl import load_workbook
        from pypdf import PdfReader
        with tempfile.TemporaryDirectory() as temporary:
            literal = "원문의 **표시**와 <b>문자</b>를 그대로 보존한다."
            plan = {"title":"가독성 계약 시험","summary":["핵심은 **근거 2건**이다."],
                    "sources":[{"id":"literal", "kind":"other", "title":"시험 원문", "text":literal}],
                    "sections":[{"title":"서식과 원문","blocks":[
                        {"type":"paragraph", "text":"결론은 **보류**이며 *추가 확인*이 필요하다. <b>태그</b>", "source_ids":["literal"]},
                        {"type":"paragraph", "text":literal, "style":"quote", "source_ids":["literal"]}]}]}
            result = reporting.render_report(plan,Path(temporary)/"reader")
            self.assertTrue(result["ready"],result)
            doc = Document(Path(result["directory"])/"02_최종 보고서"/"보고서.docx")
            paragraph = next(p for p in doc.paragraphs if p.text.startswith("결론은"))
            self.assertTrue(next(r for r in paragraph.runs if r.text == "보류").bold)
            self.assertTrue(next(r for r in paragraph.runs if r.text == "추가 확인").italic)
            self.assertIn("<b>태그</b>", paragraph.text)
            self.assertNotIn("**", paragraph.text)
            quote = next(p for p in doc.paragraphs if p.text == literal)
            self.assertEqual(quote.style.name, "Quote")
            self.assertGreater(quote.style.paragraph_format.left_indent, 0)
            self.assertEqual(len(doc.inline_shapes), 0)
            raw = load_workbook(Path(result["directory"])/"01_자료집"/"01_참고 자료.xlsx").active
            self.assertEqual(raw["E7"].value, literal)
            pdf = PdfReader(Path(result["directory"])/"02_최종 보고서"/"보고서.pdf")
            text = "".join(p.extract_text() for p in pdf.pages)
            self.assertIn("보류", text)
            self.assertNotIn("**근거 2건**", text)
            self.assertIn("**표시**", text)
            self.assertIn("<b>태그</b>", text)
            fonts = [f.get_object().get("/BaseFont", "") for page in pdf.pages for f in page["/Resources"]["/Font"].values()]
            if Path("C:/Windows/Fonts/malgunbd.ttf").is_file():
                self.assertTrue(any("Bold" in str(font) for font in fonts), fonts)

    def test_literal_citation_titles_caption_and_chart_do_not_overlap(self):
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
        from openpyxl.utils.units import EMU_to_pixels
        plan = multi_source_plan()
        with tempfile.TemporaryDirectory() as temporary:
            result = reporting.render_report(plan,Path(temporary)/"reader")
            self.assertTrue(result["ready"],result)
            data = load_workbook(Path(result["directory"])/"01_자료집"/"보고서 도표.xlsx")
            raw = load_workbook(Path(result["directory"])/"01_자료집"/"01_참고 자료.xlsx")
            self.assertEqual(raw.active["B7"].value,plan["sources"][0]["title"])
            self.assertEqual(raw.active["B7"].data_type,"s")
            for ws,block in zip(data.worksheets,plan["sections"][0]["blocks"]):
                citations = [c for row in ws for c in row if c.hyperlink]
                self.assertEqual(len(citations),5)
                self.assertEqual([c.value for c in citations],
                                 [f"출처 [{i}] · 자료집 원문" for i in range(1,6)])
                self.assertTrue(all(c.data_type == "s" for c in citations))
                self.assertTrue(all(c.data_type != "f" for row in ws for c in row))
                image = ws._images[0]
                image_row = image.anchor._from.row+1
                self.assertGreater(image_row,max(c.row for c in citations))
                caption = next(c for row in ws for c in row if c.value == block["caption"])
                self.assertGreater(caption.row,image_row+EMU_to_pixels(image.anchor.ext.cy)/20)
                self.assertTrue(caption.alignment.wrap_text)
                self.assertGreaterEqual(ws.row_dimensions[caption.row].height,26)
                if "\n" in block["caption"]:
                    self.assertGreater(ws.row_dimensions[caption.row].height,26)
                printable_bottom = range_boundaries(str(ws.print_area).split("!")[-1])[3]
                self.assertGreaterEqual(printable_bottom,caption.row)

    def test_wrapped_pdf_title_and_missing_title(self):
        from pypdf import PdfReader
        import re
        plan = example_plan()
        plan["title"] = "DUBJ2EA 후기·기존 기능 대조와 개선안 보류 보고서"
        with tempfile.TemporaryDirectory() as temporary:
            result = reporting.render_report(plan,Path(temporary)/"wrapped")
            self.assertTrue(result["ready"],result)
            text = PdfReader(Path(result["directory"])/"02_최종 보고서"/"보고서.pdf").pages[0].extract_text()
            self.assertNotIn(plan["title"],text)
            self.assertIn(re.sub(r"\s+","",plan["title"]),re.sub(r"\s+","",text))
            writer = reporting._write_pdf
            def wrong_title(doc,*args):
                changed = copy.deepcopy(doc)
                changed["title"] = "다른 문서 제목"
                return writer(changed,*args)
            with patch.object(reporting,"_write_pdf",side_effect=wrong_title):
                failure = reporting.render_report(plan,Path(temporary)/"wrong")
            self.assertFalse(failure["ready"])
            self.assertEqual(failure["checks"]["errors"][0]["code"],"reporting_verification_failed")

    def test_generation_preserves_raw_links_totals_and_no_internal_files(self):
        from openpyxl import load_workbook
        from pypdf import PdfReader
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary)/"reader"
            plan = example_plan()
            result = reporting.render_report(plan, destination)
            self.assertTrue(result["ready"], result)
            self.assertEqual(set(result), {"ready", "directory", "files", "checks"})
            self.assertTrue(result["checks"]["structure_verified"])
            self.assertEqual(result["checks"]["semantic_status"], "unverified")
            self.assertEqual(result["checks"]["visual_review"], "not_run")
            self.assertEqual(result["checks"]["chart_editing"],
                             "rendered_images; data_cells_editable; regenerate_from_report_plan")
            files = list(destination.rglob("*"))
            self.assertEqual({p.suffix for p in files if p.is_file()}, {".xlsx", ".docx", ".pdf"})
            self.assertTrue(all(set(f) == {"path", "kind", "title"} and Path(f["path"]).is_file() for f in result["files"]))
            raw = next(p for p in files if p.name.endswith("고객 후기.xlsx"))
            wb = load_workbook(raw)
            ws = wb["자료"]
            self.assertEqual("".join(ws.cell(r,5).value or "" for r in range(7,ws.max_row+1)),plan["sources"][0]["text"])
            self.assertEqual(ws["E7"].data_type, "s")
            self.assertEqual(ws["G7"].hyperlink.target, plan["sources"][0]["url"])
            self.assertEqual(ws["F7"].value, "합성")
            self.assertTrue(ws.print_area)
            data = load_workbook(destination/"01_자료집"/"보고서 도표.xlsx")
            self.assertEqual(data.worksheets[0]["C9"].border.top.style, "medium")
            self.assertEqual(data.worksheets[0]["C8"].font.color.rgb, "00B42318")
            self.assertEqual(len(data.worksheets[1]._images), 1)
            self.assertEqual([data.worksheets[1].cell(r,3).value for r in (7,8)], [4,6])
            self.assertEqual([data.worksheets[4].cell(r,4).value for r in (7,8,9)], [1,3,7])
            self.assertTrue(data.worksheets[1]["B10"].hyperlink.target.endswith("'!E7"))
            pdf = PdfReader(destination/"02_최종 보고서"/"보고서.pdf")
            self.assertTrue(pdf.pages)
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
            self.assertIn(plan["title"],text)
            self.assertIn("“문을 열 때 손잡이가 불편했다。”".replace("。","."), text)
            self.assertFalse(any(s["id"] in text for s in plan["sources"]))
            uris = [str(a.get_object().get("/A",{}).get("/URI","")) for page in pdf.pages for a in page.get("/Annots",[])]
            self.assertTrue(any("#'"+ "%EC%9E%90%EB%A3%8C" in u and "!E" in u for u in uris))
            with zipfile.ZipFile(destination/"02_최종 보고서"/"보고서.docx") as archive:
                xml = archive.read("word/document.xml").decode()
                self.assertIn("tblHeader",xml)
                self.assertIn("cantSplit",xml)
                self.assertIn("B42318",xml)
                tree = ET.fromstring(xml)
                for table in tree.findall(".//w:tbl",reporting.NS):
                    header = table.find("w:tr",reporting.NS)
                    for paragraph in header.findall(".//w:p",reporting.NS):
                        self.assertIsNotNone(paragraph.find("w:pPr/w:keepNext",reporting.NS))
                rels = archive.read("word/_rels/document.xml.rels").decode()
                self.assertIn("!E",rels)
            self.assertFalse(list(Path(temporary).glob(".reader-*")))
            with patch.object(reporting.importlib,"import_module",side_effect=ImportError("missing fixture dependency")):
                failed = reporting.render_report(example_plan(),Path(temporary)/"missing-dependency")
            self.assertFalse(failed["ready"])
            self.assertEqual(failed["checks"]["errors"][0]["code"],"reporting_dependency_missing")

    def test_existing_results_and_failure_are_truthful(self):
        with tempfile.TemporaryDirectory() as temporary:
            existing = Path(temporary)/"existing"
            existing.mkdir()
            marker = existing/"keep.txt"
            marker.write_text("keep")
            result = reporting.render_report(example_plan(),existing)
            self.assertFalse(result["ready"])
            self.assertEqual(marker.read_text(),"keep")
            self.assertEqual(result["checks"]["errors"][0]["code"],"report_destination_exists")
            for target_function in ("_deps","_write_pdf"):
                target = Path(temporary)/target_function
                with patch.object(reporting,target_function,side_effect=ToolError("test_failure","synthetic failure")):
                    result = reporting.render_report(example_plan(),target)
                self.assertFalse(result["ready"])
                self.assertEqual(result["files"],[])
                self.assertFalse(target.exists())
                self.assertEqual(result["checks"]["errors"][0]["code"],"test_failure")
            self.assertFalse(list(Path(temporary).glob(".reader-*")))

    def test_minimal_report_and_missing_metadata_are_not_absence(self):
        from openpyxl import load_workbook
        plan = {"title":"본문만 있는 시험","summary":["간단한 문단"],"sources":[
            {"id":"test-source","kind":"other","title":"원문","text":"기록됨"}],"sections":[]}
        with tempfile.TemporaryDirectory() as temporary:
            result = reporting.render_report(plan,Path(temporary)/"minimal")
            self.assertTrue(result["ready"],result)
            raw = next(Path(f["path"]) for f in result["files"] if f["kind"] == "xlsx")
            ws = load_workbook(raw).active
            self.assertEqual([ws.cell(7,i).value for i in (3,4,6,7)],["미기재"]*4)

    def test_numbered_sources_reuse_identity_preserve_shared_url_excerpts_and_exact_cells(self):
        from openpyxl import load_workbook
        from pypdf import PdfReader
        plan = {"title":"출처 연결 검증", "summary":[], "sources":[
            {"id":"first", "kind":"other", "title":"제품 문서 첫 절", "location":"첫 절", "text":"원문 가"*250, "url":"https://example.org/manual"},
            {"id":"second", "kind":"other", "title":"제품 문서 둘째 절", "location":"둘째 절", "text":"다른 절 원문", "url":"https://example.org/manual"}],
            "sections":[{"title":"분석", "blocks":[
                {"type":"paragraph", "text":"둘째 절 판단", "source_ids":["second"]},
                {"type":"paragraph", "text":"첫 절 판단", "source_ids":["first"]},
                {"type":"paragraph", "text":"같은 둘째 절 재인용", "source_ids":["second"]}]}]}
        with tempfile.TemporaryDirectory() as temporary:
            target = Path(temporary)/"reader"
            result = reporting.render_report(plan, target)
            self.assertTrue(result["ready"], result)
            self.assertEqual(result["checks"]["reference_count"], 2)
            self.assertGreaterEqual(result["checks"]["pdf_internal_links"], 3)
            with zipfile.ZipFile(target/"02_최종 보고서/보고서.docx") as z:
                root = ET.fromstring(z.read("word/document.xml"))
                anchors = [h.attrib.get("{"+reporting.NS["w"]+"}anchor") for h in root.findall(".//w:hyperlink",reporting.NS)]
                self.assertEqual([a for a in anchors if a], ["report_source_1", "report_source_2", "report_source_1"])
                bookmarks = root.findall(".//w:bookmarkStart",reporting.NS)
                self.assertEqual(len(bookmarks),2)
                titles = [e.text for e in root.findall(".//w:t",reporting.NS)]
                self.assertIn("[1] 제품 문서 둘째 절", titles)
                self.assertIn("[2] 제품 문서 첫 절", titles)
                self.assertNotIn("첫 절", titles)  # not repeated after the literal title
            raw = load_workbook(target/"01_자료집/01_참고 자료.xlsx").active
            doc = reporting._adapt_plan(plan)
            rows, mapping = reporting._source_rows(doc["sources"])
            for source in doc["sources"]:
                info = mapping[source["id"]]
                self.assertEqual("".join(raw.cell(info["row"]+i,5).value for i in range(info["chunks"])),source["text"])
            pdf = PdfReader(target/"02_최종 보고서/보고서.pdf")
            uris = [str(a.get_object().get("/A",{}).get("/URI","")) for p in pdf.pages for a in p.get("/Annots",[])]
            for source in doc["sources"]:
                cell = "E"+str(mapping[source["id"]]["row"])
                self.assertTrue(any(u.endswith("!"+cell) for u in uris))

    def test_missing_word_citations_and_pdf_citations_to_another_valid_source_are_rejected(self):
        from reportlab.platypus import Paragraph
        plan = multi_source_plan()
        with tempfile.TemporaryDirectory() as temporary:
            with patch.object(reporting, "_docx_anchor_link", return_value=None):
                missing = reporting.render_report(plan, Path(temporary)/"missing-word-links")
            self.assertFalse(missing["ready"])
            self.assertEqual(missing["checks"]["errors"][0]["code"], "reporting_verification_failed")
            def wrong_target(text, *args, **kwargs):
                if text.startswith("출처 "):
                    text = text.replace('href="#report_source_1"', 'href="#report_source_2"')
                return Paragraph(text, *args, **kwargs)
            with patch("reportlab.platypus.Paragraph", side_effect=wrong_target):
                wrong = reporting.render_report(plan, Path(temporary)/"wrong-pdf-target")
            self.assertFalse(wrong["ready"])
            self.assertEqual(wrong["checks"]["errors"][0]["code"], "reporting_verification_failed")


if __name__ == "__main__":
    unittest.main()
