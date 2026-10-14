"""Model-independent report publishing. No model, network, or office automation.

Only validated local sources and caller-authored content blocks are rendered.
Publication is a new, atomic directory; audit JSON stays outside reader folders.
"""
from __future__ import annotations

import copy
import html
import importlib
import math
import os
from pathlib import Path
import re
import shutil
import tempfile
import unicodedata
from urllib.parse import quote, urlparse
import zipfile
import xml.etree.ElementTree as ET

from .contracts import ToolError, fail

RED = "B42318"
GREY = "737373"
LINE = "D9D9D9"
FONT = "맑은 고딕"
REPORTING_PROFILE = "numbered-sources-v6"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def validate_report_plan(plan: dict) -> dict:
    """Validate the reader contract; this does not certify factual meaning."""
    def error(message):
        fail("invalid_report_plan", message)

    def fields(obj, required, optional=()):
        if not isinstance(obj, dict) or set(required)-obj.keys() or obj.keys()-set(required)-set(optional):
            error("Invalid fields; expected " + ", ".join(required))

    def text(value, label):
        if not isinstance(value, str) or not value.strip():
            error(label + " must be nonempty text")
        if any((ord(c) < 32 and c not in "\n\r\t") or 0xD800 <= ord(c) <= 0xDFFF for c in value):
            error(label + " contains unsupported XML characters")
        return value

    def array(value, label):
        if not isinstance(value, list):
            error(label + " must be an array")
        return value

    def positions(value, rows, cols, label):
        array(value, label)
        if any(not isinstance(p, list) or len(p) != 2 or any(type(n) is not int for n in p)
               or not (0 <= p[0] < rows and 0 <= p[1] < cols) for p in value):
            error(label + " contains an invalid position")
        if len({tuple(p) for p in value}) != len(value):
            error(label + " contains duplicate positions")

    fields(plan, ("title", "summary", "sources", "sections"), ("unfulfilled_requests",))
    text(plan["title"], "title")
    for paragraph in array(plan["summary"], "summary"):
        text(paragraph, "summary paragraph")
    ids = set()
    for source in array(plan["sources"], "sources"):
        fields(source, ("id", "kind", "title", "text"), ("url", "location", "model", "synthetic"))
        for key in ("id", "kind", "title", "text"):
            text(source[key], "source " + key)
        if source["id"] in ids:
            error("Duplicate source id")
        ids.add(source["id"])
        for key in ("location", "model", "url"):
            if key in source:
                text(source[key], "source " + key)
        if "synthetic" in source and type(source["synthetic"]) is not bool:
            error("source synthetic must be boolean when supplied")
        if "url" in source:
            parsed = urlparse(source["url"])
            if parsed.scheme not in ("http", "https") or not parsed.hostname or parsed.username or parsed.password:
                error("Source URL must be HTTP(S) without credentials")

    def refs(block, required=False):
        value = array(block.get("source_ids", []), "source_ids")
        if any(not isinstance(s, str) or s not in ids for s in value) or len(set(value)) != len(value):
            error("Unknown or duplicate source reference")
        if required and not value:
            error("Charts require source references")

    normalized = copy.deepcopy(plan)
    content = bool(plan["summary"])
    for section in array(normalized["sections"], "sections"):
        fields(section, ("title", "blocks"))
        text(section["title"], "section title")
        for block in array(section["blocks"], "blocks"):
            if not isinstance(block, dict):
                error("A block must be an object")
            kind = block.get("type")
            if kind == "paragraph":
                fields(block, ("type", "text"), ("source_ids", "style"))
                text(block["text"], "paragraph")
                if block.get("style", "body") not in ("body", "quote"):
                    error("Paragraph style must be body or quote")
            elif kind == "table":
                fields(block, ("type", "title", "columns", "rows"), ("source_ids", "total_rows", "highlight_cells"))
                text(block["title"], "table title")
                columns = array(block["columns"], "columns")
                if not columns:
                    error("A table needs columns")
                for column in columns:
                    text(column, "column")
                rows = array(block["rows"], "rows")
                for row in rows:
                    if not isinstance(row, list) or len(row) != len(columns):
                        error("Table row width must match columns")
                    for value in row:
                        if isinstance(value, str):
                            if any((ord(c) < 32 and c not in "\n\r\t") or 0xD800 <= ord(c) <= 0xDFFF for c in value):
                                error("Invalid table text")
                        elif not _number(value):
                            error("Table cells must be strings or finite numbers")
                totals = array(block.get("total_rows", []), "total_rows")
                if any(type(n) is not int or not 0 <= n < len(rows) for n in totals) or len(set(totals)) != len(totals):
                    error("Invalid total_rows")
                positions(block.get("highlight_cells", []), len(rows), len(columns), "highlight_cells")
            elif kind == "chart":
                fields(block, ("type", "title", "chart_type", "labels", "series", "source_ids"),
                       ("unit", "caption", "highlights"))
                text(block["title"], "chart title")
                if block["chart_type"] not in ("bar", "line", "pie", "donut", "scatter"):
                    error("Unknown chart_type")
                labels = array(block["labels"], "labels")
                series = array(block["series"], "series")
                if not labels or not series:
                    error("Charts need labels and series")
                for label in labels:
                    text(label, "chart label")
                for item in series:
                    fields(item, ("name", "values"), ("x_values",))
                    text(item["name"], "series name")
                    values = array(item["values"], "values")
                    if len(values) != len(labels) or any(not _number(n) for n in values):
                        error("Chart values must be finite and match labels")
                    if "x_values" in item:
                        xs = array(item["x_values"], "x_values")
                        if block["chart_type"] != "scatter" or len(xs) != len(labels) or any(not _number(n) for n in xs):
                            error("x_values must be finite scatter coordinates matching labels")
                    elif block["chart_type"] == "scatter":
                        error("Scatter series require x_values")
                for key in ("unit", "caption"):
                    if key in block:
                        text(block[key], key)
                positions(block.get("highlights", []), len(series), len(labels), "highlights")
                if block["chart_type"] in ("pie", "donut"):
                    values = series[0]["values"]
                    if len(series) != 1 or any(n < 0 for n in values) or not _number(sum(values)) or sum(values) <= 0:
                        error("Pie/donut requires one nonnegative complete composition")
                    if not block.get("unit") or not block.get("caption"):
                        error("Pie/donut needs a unit and caption describing the whole composition")
                refs(block, True)
            else:
                error("Unknown block type")
            refs(block)
            content = True
    for request in array(normalized.get("unfulfilled_requests", []), "unfulfilled_requests"):
        fields(request, ("request", "reason"))
        text(request["request"], "unfulfilled request")
        text(request["reason"], "unfulfilled reason")
        content = True
    if not content:
        error("Report needs content or an unfulfilled request")
    return normalized


def _adapt_plan(plan):
    """Map only validated values to private rendering structures."""
    doc = copy.deepcopy(plan)
    doc["intro"] = "\n".join(doc["summary"])
    doc["requested_gaps"] = [{"title": "수행하지 못한 요청", "request": v["request"], "reason": v["reason"]}
                             for v in doc.get("unfulfilled_requests", [])]
    for source in doc["sources"]:
        source["source_id"] = source["id"]
    # One report-wide number per exact evidence record, in first-citation order.
    # Shared URLs do not imply the same excerpt, document section or reviewer.
    source_by_id = {s["id"]: s for s in doc["sources"]}
    cited_ids = list(dict.fromkeys(sid for section in doc["sections"]
                                  for block in section["blocks"] for sid in block.get("source_ids", [])))
    doc["references"] = [source_by_id[sid] for sid in cited_ids]
    for number, source in enumerate(doc["references"], 1):
        source["reference_number"] = number
    labels = {"review": "고객 후기", "product_document": "제품 자료", "market": "시장 자료",
              "survey": "설문 자료", "interview": "인터뷰 자료", "usability": "사용성 자료", "other": "참고 자료"}
    groups = {}
    for source in doc["sources"]:
        groups.setdefault(source["kind"], []).append(source["id"])
    doc["source_groups"] = [{"title": labels.get(kind, kind), "source_ids": ids, "format": "xlsx"} for kind, ids in groups.items()]
    for section in doc["sections"]:
        for block in section["blocks"]:
            if block["type"] == "table":
                block["headers"] = block["columns"]
            elif block["type"] == "chart":
                block["kind"] = "doughnut" if block["chart_type"] == "donut" else block["chart_type"]
                block["categories"] = block["labels"]
    return doc


def _deps():
    modules = {}
    for name in ("openpyxl", "docx", "reportlab", "PIL.Image", "pypdf"):
        try:
            modules[name] = importlib.import_module(name)
        except (ImportError, OSError) as exc:
            fail("reporting_dependency_missing", "Install the project dependencies: uv sync --locked. Missing " + name + ": " + str(exc))
    return modules


def _number(value):
    try:
        return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)
    except OverflowError:
        return False


def _safe_name(value):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', " ", value).strip(" .")
    name = re.sub(r"\s+", " ", name)[:70].rstrip(" .") or "보고서"
    if re.fullmatch(r"(?i)(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])", name):
        name = "보고서 " + name
    return name



def _display(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "예" if value else "아니요"
    if isinstance(value, float):
        return format(value, ".15g")
    return str(value)


def _cell_text(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Source text beginning '=' is evidence, never a formula.


def _style_sheet(ws, title, headers, rows, deps, highlights=(), total_rows=(), widths=None, paper=None):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3.5
    ws.row_dimensions[1].height = 12
    ws.row_dimensions[2].height = 32
    ws.row_dimensions[3].height = 23
    _cell_text(ws["B2"], title)
    ws["B2"].font = Font(name=FONT, size=17, bold=True, color="000000")
    ws["B3"] = "원문과 출처"
    ws["B3"].font = Font(name=FONT, size=10, color="595959")
    if len(headers) > 1:
        ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=len(headers)+1)
    side = Side(style="thin", color=LINE)
    highlight = set(tuple(p) for p in highlights)
    for col, heading in enumerate(headers, 2):
        cell = ws.cell(6, col)
        _cell_text(cell, heading)
        cell.font = Font(name=FONT, size=10, bold=True, color="000000")
        cell.fill = PatternFill("solid", fgColor="F2F2F2")
        cell.border = Border(top=side, bottom=side, left=side, right=side)
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
        longest = max([len(heading)+5]+[min(60, len(_display(r[col-2]))+3) for r in rows])
        ws.column_dimensions[get_column_letter(col)].width = widths[col-2] if widths else min(60, max(14, longest))
    ws.row_dimensions[6].height = 28
    for ridx, row in enumerate(rows):
        lines = 1
        for cidx, value in enumerate(row):
            cell = ws.cell(ridx+7, cidx+2)
            _cell_text(cell, value)
            numeric = _number(value)
            year = numeric and "년" in headers[cidx] or headers[cidx].lower() == "year"
            cell.alignment = Alignment(horizontal="center" if year else "right" if numeric else "left", vertical="center", indent=0 if year else 1, wrap_text=True)
            cell.font = Font(name=FONT, size=10, bold=ridx in total_rows, color=RED if (ridx, cidx) in highlight else "000000")
            cell.fill = PatternFill("solid", fgColor="FFFFFF")
            cell.border = Border(top=Side(style="medium" if ridx in total_rows else "thin", color=LINE),
                                 bottom=side, left=side, right=side)
            if numeric:
                cell.number_format = "#,##0" if value == int(value) else "#,##0.0##############"
            width = ws.column_dimensions[get_column_letter(cidx+2)].width
            # Excel column widths use narrow-character units. Korean/other wide
            # glyphs occupy about two; explicit line breaks count separately.
            # Extra vertical space protects Malgun ascenders in printed cells.
            lines = max(lines, sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(char) in ("W", "F") else 1
                for char in line)/max(1, width-3))) for line in _display(value).split("\n")))
        # Excel supports at most 409 points. Long source chunks continue below.
        if lines*16+14 > 409:
            fail("reporting_layout_failed", "An Excel row exceeds its readable height; shorten table cells or source metadata")
        ws.row_dimensions[ridx+7].height = max(30, lines*16+14)
    ws.column_dimensions[get_column_letter(len(headers)+2)].width = 3.5
    ws.freeze_panes = "B7"
    if rows:
        ws.auto_filter.ref = f"B6:{get_column_letter(len(headers)+1)}{len(rows)+6}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = paper or ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = "1:6"
    ws.print_options.horizontalCentered = True
    ws.print_area = f"B2:{get_column_letter(len(headers)+1)}{max(6,len(rows)+6)}"
    from openpyxl.worksheet.page import PageMargins
    ws.page_margins = PageMargins(left=.3, right=.3, top=.4, bottom=.4, header=.2, footer=.2)
    _sheet_title(ws, title, len(headers)+1)
    return ws


def _sheet_title(ws, title, end_column):
    """Fit the full heading to its actual table/chart print width."""
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
    for region in list(ws.merged_cells.ranges):
        if region.min_row == region.max_row == 2:
            ws.unmerge_cells(str(region))
    if end_column > 2:
        ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=end_column)
    ws["B2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    width = sum(ws.column_dimensions[get_column_letter(col)].width for col in range(2, end_column+1))
    lines = sum(max(1, math.ceil(len(line)*3.1/max(8, width-3))) for line in title.split("\n"))
    height = max(32, lines*24+8)
    if height > 409:
        fail("reporting_layout_failed", "An Excel title exceeds its readable row height")
    ws.row_dimensions[2].height = height


def _source_rows(records):
    rows, first_rows = [], {}
    for source in records:
        # Short continued rows preserve text and stay readable in Excel.
        chunks, remaining = [], source["text"]
        while remaining:
            take = min(480,len(remaining))
            newlines = [m.end() for m in re.finditer("\n",remaining[:take])]
            if len(newlines) > 8:
                take = newlines[7]
            chunks.append(remaining[:take])
            remaining = remaining[take:]
        first_rows[source["source_id"]] = {"row": len(rows)+7, "chunks":len(chunks)}
        for idx, chunk in enumerate(chunks):
            rows.append([source["title"] if not idx else source["title"]+" 계속",
                         source.get("model", "미기재"), source.get("location", "미기재"), chunk,
                         "합성" if source.get("synthetic") is True else "원자료" if source.get("synthetic") is False else "미기재",
                         source.get("url", "미기재")])
    return rows, first_rows


def _write_libraries(doc, folder, deps):
    from openpyxl import Workbook
    records = {s["source_id"]: s for s in doc["sources"]}
    source_links, outputs = {}, []
    for index, group in enumerate(doc["source_groups"], 1):
        group_records = [records[s] for s in group["source_ids"]]
        stem = f"{index:02d}_"+_safe_name(group["title"])
        path = folder/(stem+".xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "자료"
        rows, mapping = _source_rows(group_records)
        _style_sheet(ws, group["title"], ["자료명", "모델과 버전", "원문 위치", "원문", "자료 구분", "출처 URL"], rows, deps,
                     widths=[30, 36, 24, 64, 14, 36], paper=ws.PAPERSIZE_A3)
        ws["B3"] = "원문과 출처 · 인쇄 A3 가로 · 화면에서 확대 가능"
        for idx, row in enumerate(rows, 7):
            if row[-1] != "미기재":
                ws.cell(idx, 7).hyperlink = row[-1]
        wb.save(path)
        outputs.append(path)
        for sid, info in mapping.items():
            source_links.setdefault(sid, {"file": path.name, "sheet": "자료", "cell": "E"+str(info["row"]),
                                        "chunks":info["chunks"]})
        if group["format"] in ("docx", "both"):
            source_doc = _new_docx()
            source_doc.add_paragraph(group["title"], "Title")
            for source in group_records:
                source_doc.add_heading(source["title"], level=1)
                source_doc.add_paragraph(source["location"])
                source_doc.add_paragraph(source["text"])
                if source.get("summary"):
                    source_doc.add_paragraph(source["summary"])
                if source.get("url"):
                    _docx_link(source_doc.add_paragraph(), source["url"], source["url"])
            source_path = folder/(stem+".docx")
            source_doc.save(source_path)
            outputs.append(source_path)
    # Editable data cells and the same rendered charts as DOCX/PDF.
    tables = [(section["title"], b) for section in doc["sections"] for b in section["blocks"] if b["type"] in ("table", "chart")]
    if tables:
        wb = Workbook()
        wb.remove(wb.active)
        printed_sources = []
        for index, (title, block) in enumerate(tables, 1):
            ws = wb.create_sheet(f"자료{index}")
            if block["type"] == "table":
                _style_sheet(ws, block["title"], block["headers"], block["rows"], deps,
                             block.get("highlight_cells", []), block.get("total_rows", []))
            else:
                rows = [[cat]+[s["values"][i] for s in block["series"]] for i, cat in enumerate(block["categories"])]
                headers = ["항목"]+[s["name"] for s in block["series"]]
                if block["kind"] == "scatter":
                    headers += ["X "+s["name"] for s in block["series"]]
                    for i,row in enumerate(rows):
                        row.extend(s["x_values"][i] for s in block["series"])
                _style_sheet(ws, block["title"], headers, rows, deps,
                             [[point,series+1] for series,point in block.get("highlights",[])])
                _cell_text(ws["B3"], block.get("unit", ""))
            citation_row = len(block["rows"] if block["type"] == "table" else rows)+8
            for sid in block.get("source_ids", []):
                source = records[sid]
                info = source_links[sid]
                cell = ws.cell(citation_row,2)
                _cell_text(cell, f"출처 [{source['reference_number']}] · 자료집 원문")
                cell.hyperlink = info["file"]+"#'"+info["sheet"]+"'!"+info["cell"]
                printed_sources.append([source["reference_number"], ws.title, cell.hyperlink.target])
                _sheet_note(ws,citation_row,cell.value,13 if block["type"] == "chart" else len(block["headers"])+1)
                citation_row += 1
            if block["type"] == "chart":
                from io import BytesIO
                from openpyxl.drawing.image import Image as ExcelImage
                chart_row = citation_row+2
                for drawing in _chart_drawings(block, _font()):
                    buffer = BytesIO()
                    _drawing_png(drawing,buffer)
                    buffer.seek(0)
                    picture = ExcelImage(buffer)
                    picture.width = 720
                    picture.height = 720*drawing.height/drawing.width
                    ws.add_image(picture,"B"+str(chart_row))
                    chart_row += math.ceil(picture.height/20)+2
                if block.get("caption"):
                    _sheet_note(ws,chart_row,block["caption"],13)
                    chart_row += 1
                ws.page_setup.fitToHeight = 1
                ws.print_area = f"B2:M{chart_row+2}"
                _sheet_title(ws, block["title"], 13)
            elif block.get("source_ids"):
                from openpyxl.utils import get_column_letter
                ws.print_area = f"B2:{get_column_letter(len(block['headers'])+1)}{len(block['rows'])+6}"
        if printed_sources:
            ws = wb.create_sheet("도표 출처")
            rows = [[number, title, "자료집 원문"] for number, title, target in printed_sources]
            _style_sheet(ws, "보고서 도표 출처", ["출처 번호", "도표 시트", "원문 연결"], rows, deps,
                         widths=[14, 24, 40])
            for row, (_, _, target) in enumerate(printed_sources, 7):
                ws.cell(row, 4).hyperlink = target
        path = folder/"보고서 도표.xlsx"
        wb.save(path)
        outputs.append(path)
    return outputs, source_links


def _sheet_note(ws, row, value, end_column):
    """Keep citation titles literal and captions readable inside the print area."""
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    _cell_text(ws.cell(row,2),value)
    if end_column > 2:
        ws.merge_cells(start_row=row,start_column=2,end_row=row,end_column=end_column)
    cell = ws.cell(row,2)
    cell.font = Font(name=FONT,size=10,color="000000")
    cell.alignment = Alignment(horizontal="left",vertical="top",wrap_text=True)
    width = sum(ws.column_dimensions[get_column_letter(col)].width for col in range(2,end_column+1))
    lines = sum(max(1,math.ceil(len(line)*2.1/max(8,width-3))) for line in value.split("\n"))
    if lines*15+10 > 409:
        fail("reporting_layout_failed","An Excel caption or citation exceeds its readable row height")
    ws.row_dimensions[row].height = max(26,lines*15+10)


def _inline_runs(value):
    """A small presentation-only subset; source records remain literal."""
    text = str(value)
    at = 0
    for match in re.finditer(r"\*\*[^*\n]+\*\*|(?<!\*)\*[^*\n]+\*(?!\*)", text):
        if match.start() > at:
            yield text[at:match.start()], False, False
        token = match.group()
        bold = token.startswith("**")
        yield token[2:-2] if bold else token[1:-1], bold, not bold
        at = match.end()
    if at < len(text):
        yield text[at:], False, False


def _inline_html(value):
    parts = []
    for text, bold, italic in _inline_runs(value):
        escaped = html.escape(text).replace("\n", "<br/>")
        tag = "b" if bold else "i" if italic else None
        parts.append(f"<{tag}>{escaped}</{tag}>" if tag else escaped)
    return "".join(parts)


def _docx_paragraph(document, value, style=None, formatted=True):
    paragraph = document.add_paragraph(style=style)
    runs = _inline_runs(value) if formatted else [(str(value), False, False)]
    for text, bold, italic in runs:
        run = paragraph.add_run(text)
        if bold:
            run.bold = True
        if italic:
            run.italic = True
    return paragraph


def _new_docx():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    section.left_margin = section.right_margin = Inches(.8)
    section.top_margin = section.bottom_margin = Inches(.8)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Quote"):
        style = document.styles[name]
        style.font.name = FONT
        style.font.size = Pt(11 if name in ("Normal", "Quote") else 25 if name == "Title" else 16 if name == "Heading 1" else 12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name in ("Title", "Heading 1", "Heading 2")
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is not None:
            fonts.set(qn("w:eastAsia"), FONT)
        ppr = style.element.get_or_add_pPr()
        for tag in ("autoSpaceDE", "autoSpaceDN"):
            value = OxmlElement("w:"+tag)
            value.set(qn("w:val"), "0")
            ppr.append(value)
        for border in list(ppr.findall(qn("w:pBdr"))):
            ppr.remove(border)
        style.paragraph_format.space_after = Pt(8)
    document.styles["Normal"].paragraph_format.line_spacing = 1.2
    document.styles["Title"].paragraph_format.keep_with_next = True
    for name in ("Heading 1", "Heading 2"):
        document.styles[name].paragraph_format.space_before = Pt(16)
    quote = document.styles["Quote"]
    quote.font.italic = True
    quote.font.color.rgb = RGBColor.from_string("595959")
    quote.paragraph_format.left_indent = Inches(.22)
    quote.paragraph_format.right_indent = Inches(.12)
    quote.paragraph_format.line_spacing = 1.2
    return document


def _docx_link(paragraph, label, target):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _source_targets(ids, sources, links):
    for sid in ids:
        source = sources[sid]
        link = links[sid]
        local = "../01_자료집/"+link["file"]+"#'"+link["sheet"]+"'!"+link["cell"]
        yield source, quote(local, safe="/!#'"), link["sheet"]+" "+link["cell"]


def _docx_sources(document, ids, sources, links):
    if not ids:
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = 10
    paragraph.add_run("출처 ")
    for idx, sid in enumerate(ids):
        if idx:
            paragraph.add_run(", ")
        number = sources[sid]["reference_number"]
        _docx_anchor_link(paragraph, f"[{number}]", f"report_source_{number}")
    for run in paragraph.runs:
        from docx.shared import Pt
        run.font.size = Pt(9)


def _docx_anchor_link(paragraph, label, anchor):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    for tag, value in (("w:color", GREY), ("w:sz", "18")):
        prop = OxmlElement(tag)
        prop.set(qn("w:val"), value)
        props.append(prop)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _reference_notes(source):
    # Full provenance stays literal in the raw workbook. Display only the
    # model/document/time labels when our canonical normalizer supplies them.
    model = source.get("model", "")
    parts = model.split(" / ")
    prefixes = ("모델 ", "문서 버전 ", "관찰일: ", "게시일: ")
    selected = [p for p in parts if p.startswith(prefixes) and not p.startswith("모델 버전 ")]
    return " · ".join(selected) if selected else model


def _reference_location(source):
    location = source.get("location", "")
    return location if location and location not in source["title"] else ""


def _docx_references(document, doc, links):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    if not doc["references"]:
        return
    document.add_heading("출처 목록", level=1)
    document.add_paragraph("본문의 출처 번호를 누르면 해당 자료로 이동합니다. 전체 원문과 상세 수집 정보는 자료집의 연결된 셀에서 확인할 수 있습니다.")
    for source in doc["references"]:
        number = source["reference_number"]
        paragraph = document.add_paragraph(f"[{number}] {source['title']}")
        paragraph.paragraph_format.keep_with_next = True
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), str(number))
        start.set(qn("w:name"), f"report_source_{number}")
        end.set(qn("w:id"), str(number))
        paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
        paragraph._p.append(end)
        for detail in (_reference_location(source), _reference_notes(source)):
            if detail:
                note = document.add_paragraph(detail)
                note.paragraph_format.keep_with_next = True
                for run in note.runs:
                    run.font.size = Pt(9)
        paragraph = document.add_paragraph()
        if source.get("url"):
            _docx_link(paragraph, "웹 원문", source["url"])
            paragraph.add_run(" · ")
        _, local, cell = next(_source_targets([source["id"]], {source["id"]: source}, links))
        _docx_link(paragraph, "자료집 원문 ("+cell+")", local)


def _column_weights(headers, rows):
    weights = [max(5, len(h)*.7, max((min(40, len(_display(r[i])))**.5*2 for r in rows), default=5)) for i, h in enumerate(headers)]
    total = sum(weights)
    return [v/total for v in weights]


def _docx_table(document, headers, rows, highlights=(), total_rows=()):
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    widths = _column_weights(headers, rows)
    for col, weight in zip(table.columns, widths):
        col.width = Inches(6.6677*weight)
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = _display(value)
            for p in cells[idx].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if _number(value) else WD_ALIGN_PARAGRAPH.LEFT
    for ridx, row in enumerate(table.rows):
        rowpr = row._tr.get_or_add_trPr()
        if not any(len(_display(v)) > 800 for v in (headers if ridx == 0 else rows[ridx-1])):
            rowpr.append(OxmlElement("w:cantSplit"))
        for cidx, cell in enumerate(row.cells):
            cell.width = Inches(6.6677*widths[cidx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for name in ("top", "left", "bottom", "right"):
                border = OxmlElement("w:"+name)
                for key, value in (("val", "single"), ("sz", "12" if name == "top" and ridx-1 in total_rows else "4"), ("color", LINE)):
                    border.set(qn("w:"+key), value)
                borders.append(border)
            tcpr.append(borders)
            if ridx == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "F2F2F2")
                tcpr.append(shade)
            margins = OxmlElement("w:tcMar")
            for name, value in (("top", 90), ("bottom", 90), ("left", 100), ("right", 100)):
                margin = OxmlElement("w:"+name)
                margin.set(qn("w:w"), str(value))
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            tcpr.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.keep_with_next = ridx == 0
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.bold = ridx == 0 or ridx-1 in total_rows
                    run.font.color.rgb = RGBColor.from_string(RED if [ridx-1,cidx] in highlights else "000000")
    document.add_paragraph()
    return table


def _font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    if "ReportKorean" in pdfmetrics.getRegisteredFontNames():
        return "ReportKorean"
    candidates = [Path(os.environ.get("WINDIR", "C:/Windows"))/"Fonts/malgun.ttf", Path("/usr/share/fonts/truetype/noto/NotoSansKR-Regular.ttf")]
    for path in candidates:
        if path.is_file():
            try:
                pdfmetrics.registerFont(TTFont("ReportKorean", str(path)))
                bold_path = path.with_name("malgunbd.ttf" if path.name == "malgun.ttf" else "NotoSansKR-Bold.ttf")
                bold = "ReportKorean"
                if bold_path.is_file():
                    pdfmetrics.registerFont(TTFont("ReportKoreanBold", str(bold_path)))
                    bold = "ReportKoreanBold"
                pdfmetrics.registerFontFamily("ReportKorean", normal="ReportKorean", bold=bold,
                                               italic="ReportKorean", boldItalic=bold)
                return "ReportKorean"
            except Exception:
                continue
    pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
    pdfmetrics.registerFontFamily("HYSMyeongJo-Medium", normal="HYSMyeongJo-Medium",
                                  bold="HYSMyeongJo-Medium", italic="HYSMyeongJo-Medium",
                                  boldItalic="HYSMyeongJo-Medium")
    return "HYSMyeongJo-Medium"


def _chart_drawings(block, font):
    from reportlab.graphics.shapes import Drawing, Rect, String, Line, Circle, Wedge
    from reportlab.lib.colors import HexColor, black, white
    from reportlab.pdfbase.pdfmetrics import stringWidth
    grid, grey, red = HexColor("#E6E6E6"), HexColor("#"+GREY), HexColor("#"+RED)
    highlights = {tuple(p) for p in block.get("highlights", [])}
    cats, series, kind = block["categories"], block["series"], block["kind"]
    palette = ["737373", "303030", "A6A6A6", "595959", "BDBDBD", "454545"]
    dashes = [None, [6,3], [2,3], [6,2,1,2]]

    def series_color(index):
        return HexColor("#"+palette[index % len(palette)])

    def marker(d, x, y, index, color):
        if index % 2:
            d.add(Rect(x-3.5,y-3.5,7,7,fillColor=color,strokeColor=white))
        else:
            d.add(Circle(x,y,3.5,fillColor=color,strokeColor=white))

    def label(d, x, y, value, width=120, size=9, color=black, max_lines=3):
        value = str(value)
        lines, current = [], ""
        for char in value:
            if stringWidth(current+char,font,size) > width and current:
                lines.append(current)
                current = char
            else:
                current += char
        lines.append(current)
        if len(lines) > max_lines:
            fail("reporting_layout_failed", "A chart label exceeds its readable area; shorten the label")
        for index,line in enumerate(lines):
            d.add(String(x,y-index*(size+2),line,fontName=font,fontSize=size,fillColor=color))
        return len(lines)*(size+2)

    if kind in ("pie", "doughnut"):
        drawings = []
        values, total = series[0]["values"], sum(series[0]["values"])
        for start in range(0,len(cats),10):
            d = Drawing(470, max(220,min(10,len(cats)-start)*30+30))
            angle = 0
            for index,value in enumerate(values):
                color = red if (0,index) in highlights else HexColor("#"+format(100+index%5*26,"02X")*3)
                if value:
                    d.add(Wedge(95,d.height/2,75,angle,angle+value/total*360,fillColor=color,strokeColor=white))
                angle += value/total*360
            if kind == "doughnut":
                d.add(Circle(95,d.height/2,40,fillColor=white,strokeColor=white))
            for index in range(start,min(len(cats),start+10)):
                y = d.height-25-(index-start)*30
                color = red if (0,index) in highlights else HexColor("#"+format(100+index%5*26,"02X")*3)
                d.add(Rect(190,y-2,8,8,fillColor=color,strokeColor=None))
                label(d,205,y,cats[index]+"  "+_display(values[index]),width=250,max_lines=2)
            drawings.append(d)
        return drawings
    allvalues = [v for item in series for v in item["values"]]
    lo, hi = min([0]+allvalues), max([0]+allvalues)
    span = hi-lo or 1
    if not _number(span):
        fail("invalid_report_plan", "Chart range overflows")
    if kind == "bar":
        drawings = []
        # Full labels get their own wrapped row; dense charts continue as figures.
        for start in range(0,len(cats),max(1,8//len(series))):
            count = min(max(1,8//len(series)),len(cats)-start)
            d = Drawing(470,90+count*46*len(series))
            x0, width, top = 150, 270, d.height-25
            for tick in range(5):
                value = lo+span*tick/4
                x = x0+width*tick/4
                d.add(Line(x,52,x,top,strokeColor=grid))
                label(d,x-12,36,format(value,".6g"),width=65,size=8)
            zero = x0+(0-lo)/span*width
            d.add(Line(x0,52,x0+width,52,strokeColor=grey))
            d.add(Line(zero,52,zero,top,strokeColor=grey))
            for index in range(count):
                absolute = start+index
                for sidx,item in enumerate(series):
                    y = top-24-(index*len(series)+sidx)*46
                    value = item["values"][absolute]
                    color = red if (sidx,absolute) in highlights else series_color(sidx)
                    target = x0+(value-lo)/span*width
                    label(d,0,y+9,cats[absolute]+(" "+item["name"] if len(series)>1 else ""),width=140,size=8)
                    d.add(Rect(min(zero,target),y,abs(target-zero),13,fillColor=color,strokeColor=None))
                    label(d,430,y+2,_display(value),width=40,size=8,color=color)
            label(d,150,15,block.get("unit","값"),width=300,size=9)
            drawings.append(d)
        return drawings
    # Categorical line and numeric scatter use actual x/y axes and pale grids.
    drawings = []
    batches = [(0,len(cats))] if kind == "scatter" else [(i,min(8,len(cats)-i)) for i in range(0,len(cats),8)]
    xs = [v for item in series for v in item["x_values"]] if kind == "scatter" else list(range(len(cats)))
    xmin,xmax = min(xs),max(xs)
    xspan = xmax-xmin or 1
    if not _number(xspan):
        fail("invalid_report_plan", "Chart x range overflows")
    for start,count in batches:
        d = Drawing(470,310+max(0,math.ceil(len(series)/3)-1)*15)
        left,bottom,width,height = 55,65,365,185
        d.add(Rect(left,bottom,width,height,fillColor=white,strokeColor=grey))
        for tick in range(5):
            value = lo+span*tick/4
            y = bottom+height*tick/4
            d.add(Line(left,y,left+width,y,strokeColor=grid))
            label(d,0,y-3,format(value,".6g"),width=50,size=8)
            if kind == "scatter":
                x = left+width*tick/4
                d.add(Line(x,bottom,x,bottom+height,strokeColor=grid))
                label(d,x-12,bottom-18,format(xmin+xspan*tick/4,".6g"),width=60,size=8)
        for sidx,item in enumerate(series):
            prior = None
            for absolute in range(start,start+count):
                x = left+(item["x_values"][absolute]-xmin)/xspan*width if kind == "scatter" else left+(absolute-start)/max(1,count-1)*width
                y = bottom+(item["values"][absolute]-lo)/span*height
                color = red if (sidx,absolute) in highlights else series_color(sidx)
                if prior is not None and kind == "line":
                    d.add(Line(prior[0],prior[1],x,y,strokeColor=series_color(sidx),
                               strokeDashArray=dashes[sidx % len(dashes)]))
                marker(d,x,y,sidx,color)
                prior = (x,y)
                if kind == "line" and sidx == 0:
                    label(d,x-18,bottom-18,cats[absolute],width=max(40,width/max(1,count)),size=8)
            legend_x, legend_y = left+(sidx%3)*130, d.height-25-(sidx//3)*15
            if kind == "line":
                d.add(Line(legend_x,legend_y+3,legend_x+20,legend_y+3,
                           strokeColor=series_color(sidx),strokeDashArray=dashes[sidx % len(dashes)]))
            marker(d,legend_x+10,legend_y+3,sidx,series_color(sidx))
            label(d,legend_x+25,legend_y,item["name"],width=100,size=9,max_lines=1)
        label(d,left,265,block.get("unit","값"),width=350,size=9)
        drawings.append(d)
    return drawings


def _write_docx(doc, target, sources, links, temp):
    from docx.shared import Inches
    document = _new_docx()
    document.add_paragraph(doc["title"], "Title")
    for paragraph in doc.get("summary", []):
        _docx_paragraph(document, paragraph)
    image_index = 0
    for section in doc["sections"]:
        document.add_heading(section["title"], level=1)
        for block in section["blocks"]:
            if block["type"] == "paragraph":
                is_quote = block.get("style") == "quote"
                paragraph = _docx_paragraph(document, block["text"], "Quote" if is_quote else None, formatted=not is_quote)
                paragraph.paragraph_format.keep_with_next = bool(block.get("source_ids"))
            elif block["type"] == "table":
                title = document.add_paragraph(block["title"], "Heading 2")
                title.paragraph_format.keep_with_next = True
                table = _docx_table(document, block["headers"], block["rows"], block.get("highlight_cells", []), block.get("total_rows", []))
                if block.get("source_ids"):
                    for cell in table.rows[-1].cells:
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.keep_with_next = True
            else:
                heading = document.add_paragraph(block["title"], "Heading 2")
                heading.paragraph_format.keep_with_next = True
                for drawing in _chart_drawings(block, _font()):
                    image_index += 1
                    image = temp/("chart"+str(image_index)+".png")
                    _drawing_png(drawing, image)
                    picture = document.add_picture(str(image), width=Inches(6.4))
                    document.paragraphs[-1].paragraph_format.keep_with_next = True
                if block.get("unit"):
                    paragraph = document.add_paragraph("단위 "+block["unit"])
                    paragraph.paragraph_format.keep_with_next = True
                if block.get("caption"):
                    paragraph = document.add_paragraph(block["caption"])
                    paragraph.paragraph_format.keep_with_next = bool(block.get("source_ids"))
            _docx_sources(document, block.get("source_ids", []), sources, links)
    for gap in doc["requested_gaps"]:
        document.add_heading(gap["title"],level=1)
        document.add_paragraph(gap["request"])
        document.add_paragraph(gap["reason"])
    _docx_references(document, doc, links)
    document.save(target)


def _drawing_png(drawing, target):
    # ReportLab's Cairo/renderPM extension is not available in every install.
    # Rasterize its own vector chart primitives with Pillow, preserving data.
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.graphics.shapes import Rect, String, Line, Circle, Wedge
    scale = 3
    image = Image.new("RGB", (int(drawing.width*scale),int(drawing.height*scale)), "white")
    paint = ImageDraw.Draw(image)
    font_path = Path(os.environ.get("WINDIR", "C:/Windows"))/"Fonts/malgun.ttf"
    def color(c):
        return tuple(round(v*255) for v in (c.red,c.green,c.blue)) if c is not None else None
    def point(x,y):
        return x*scale,(drawing.height-y)*scale
    for item in drawing.contents:
        if isinstance(item,String):
            font = ImageFont.truetype(str(font_path), round(item.fontSize*scale)) if font_path.is_file() else ImageFont.load_default(size=round(item.fontSize*scale))
            paint.text(point(item.x,item.y),item.text,font=font,fill=color(item.fillColor),anchor="ls")
        elif isinstance(item,Line):
            first,last = point(item.x1,item.y1),point(item.x2,item.y2)
            dash = item.strokeDashArray
            if dash:
                dx,dy = last[0]-first[0],last[1]-first[1]
                length = math.hypot(dx,dy)
                at,index = 0,0
                while at < length:
                    end = min(length,at+dash[index % len(dash)]*scale)
                    if index % 2 == 0 and length:
                        paint.line([(first[0]+dx*at/length,first[1]+dy*at/length),
                                    (first[0]+dx*end/length,first[1]+dy*end/length)],
                                   fill=color(item.strokeColor),width=2)
                    at,index = end,index+1
            else:
                paint.line([first,last],fill=color(item.strokeColor),width=2)
        elif isinstance(item,Rect):
            paint.rectangle([point(item.x,item.y+item.height),point(item.x+item.width,item.y)],fill=color(item.fillColor),outline=color(item.strokeColor))
        elif isinstance(item,Circle):
            paint.ellipse([point(item.cx-item.r,item.cy+item.r),point(item.cx+item.r,item.cy-item.r)],fill=color(item.fillColor),outline=color(item.strokeColor))
        elif isinstance(item,Wedge):
            paint.pieslice([point(item.centerx-item.radius,item.centery+item.radius),point(item.centerx+item.radius,item.centery-item.radius)],start=-item.endangledegrees,end=-item.startangledegrees,fill=color(item.fillColor),outline=color(item.strokeColor))
    image.save(target,format="PNG")


def _write_pdf(doc, target, sources, links):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, KeepTogether
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.colors import HexColor, black, white
    from reportlab.lib.enums import TA_RIGHT
    font = _font()
    from reportlab.pdfbase import pdfmetrics
    bold_font = "ReportKoreanBold" if "ReportKoreanBold" in pdfmetrics.getRegisteredFontNames() else font
    normal = ParagraphStyle("Report",fontName=font,fontSize=11,leading=16,spaceAfter=9,wordWrap="CJK",allowWidows=0,allowOrphans=0)
    title = ParagraphStyle("Title",parent=normal,fontName=bold_font,fontSize=25,leading=34,spaceAfter=18)
    heading = ParagraphStyle("Section",parent=normal,fontName=bold_font,fontSize=16,leading=23,spaceBefore=16,spaceAfter=10,keepWithNext=True)
    subheading = ParagraphStyle("Subheading",parent=heading,fontSize=12,leading=17,spaceBefore=9,spaceAfter=7)
    quote_style = ParagraphStyle("Quote",parent=normal,leftIndent=16,rightIndent=9,textColor=HexColor("#595959"),spaceBefore=4,spaceAfter=11)
    cellstyle = ParagraphStyle("Cell",parent=normal,fontSize=9,leading=13,spaceAfter=0)
    headerstyle = ParagraphStyle("HeaderCell",parent=cellstyle,fontName=bold_font)
    source_style = ParagraphStyle("Source",parent=normal,fontSize=8.5,leading=12,spaceAfter=13)
    flow = []
    def p(text,style=normal,formatted=False):
        return Paragraph(_inline_html(text) if formatted else html.escape(str(text)).replace("\n","<br/>"),style)
    def table(headers,rows,highlights=(),total_rows=()):
        data = [[p(h,headerstyle) for h in headers]]+[[p(_display(v),ParagraphStyle("n",parent=cellstyle,alignment=TA_RIGHT) if _number(v) else cellstyle) for v in row] for row in rows]
        widths = [v*(A4[0]-115.2) for v in _column_weights(headers,rows)]
        commands = [("GRID",(0,0),(-1,-1),.35,HexColor("#"+LINE)),("BACKGROUND",(0,0),(-1,-1),white),("TEXTCOLOR",(0,0),(-1,-1),black),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7),("LEFTPADDING",(0,0),(-1,-1),7),("RIGHTPADDING",(0,0),(-1,-1),7)]
        commands.append(("BACKGROUND",(0,0),(-1,0),HexColor("#F2F2F2")))
        for total in total_rows:
            commands.append(("LINEABOVE",(0,total+1),(-1,total+1),1,HexColor("#"+GREY)))
        for row,col in highlights:
            # Paragraph color, not only table TEXTCOLOR, controls paragraph text.
            data[row+1][col] = p(_display(rows[row][col]),ParagraphStyle("red",parent=cellstyle,
                                alignment=TA_RIGHT if _number(rows[row][col]) else cellstyle.alignment,
                                textColor=HexColor("#"+RED)))
        result = LongTable(data,colWidths=widths,repeatRows=1,splitByRow=1,splitInRow=1)
        result.setStyle(TableStyle(commands))
        flow.extend((result,Spacer(1,10)))
    def citations(ids):
        labels = []
        for sid in ids:
            number = sources[sid]["reference_number"]
            labels.append(f'<link href="#report_source_{number}">[{number}]</link>')
        if labels:
            flow.append(Paragraph("출처 "+", ".join(labels),source_style))
    flow.append(p(doc["title"],title))
    for paragraph in doc.get("summary", []):
        flow.append(p(paragraph, formatted=True))
    for section in doc["sections"]:
        section_start = len(flow)
        flow.append(p(section["title"],heading))
        for block_index, block in enumerate(section["blocks"]):
            block_start = len(flow)
            keep_figure = False
            if block["type"] == "paragraph":
                is_quote = block.get("style") == "quote"
                flow.append(p(block["text"],quote_style if is_quote else normal,formatted=not is_quote))
            elif block["type"] == "table":
                flow.append(p(block["title"],subheading))
                table(block["headers"],block["rows"],block.get("highlight_cells",[]),block.get("total_rows",[]))
            else:
                flow.append(p(block["title"],subheading))
                drawings = _chart_drawings(block,font)
                keep_figure = len(drawings) == 1
                flow.extend(drawings)
                if block.get("unit"):
                    flow.append(p("단위 "+block["unit"]))
                if block.get("caption"):
                    flow.append(p(block["caption"]))
            if block.get("source_ids") and flow and hasattr(flow[-1], "getPlainText"):
                flow[-1].keepWithNext = True
            citations(block.get("source_ids",[]))
            if keep_figure:
                if block_index == 0:
                    block_start = section_start
                figure = flow[block_start:]
                flow[block_start:] = [KeepTogether(figure)]
    for gap in doc["requested_gaps"]:
        flow.extend((p(gap["title"],heading),p(gap["request"]),p(gap["reason"])))
    if doc["references"]:
        flow.append(p("출처 목록", heading))
        flow.append(p("본문의 출처 번호를 누르면 해당 자료로 이동합니다. 전체 원문과 상세 수집 정보는 자료집의 연결된 셀에서 확인할 수 있습니다.", source_style))
        for source in doc["references"]:
            number = source["reference_number"]
            label = f'<a name="report_source_{number}"/>[{number}] '+html.escape(source["title"])
            reference = Paragraph(label, ParagraphStyle("Reference", parent=normal, fontSize=10, leading=14, keepWithNext=True))
            reference.reference_number = number
            reference.reference_title = f"[{number}] "+source["title"]
            flow.append(reference)
            for detail in (_reference_location(source), _reference_notes(source)):
                if detail:
                    flow.append(p(detail, ParagraphStyle("ReferenceNote", parent=source_style, spaceAfter=3, keepWithNext=True)))
            _, local, cell = next(_source_targets([source["id"]], sources, links))
            labels = []
            if source.get("url"):
                labels.append('<link href="'+html.escape(quote(source["url"],safe=":/?=&%#"),quote=True)+'">웹 원문</link>')
            labels.append('<link href="'+html.escape(local,quote=True)+'">자료집 원문 ('+html.escape(cell)+')</link>')
            flow.append(Paragraph(" · ".join(labels), source_style))
    def footer(canvas,pdfdoc):
        canvas.setFont(font,8)
        canvas.drawRightString(A4[0]-57.6,33,str(pdfdoc.page))
    class ReaderPDF(SimpleDocTemplate):
        def afterFlowable(self, item):
            number = getattr(item, "reference_number", None)
            if number is not None:
                key = f"report_source_{number}"
                if number == 1:
                    self.canv.addOutlineEntry("출처 목록", key, level=0, closed=False)
                self.canv.addOutlineEntry(item.reference_title, key, level=1, closed=False)
    ReaderPDF(str(target),pagesize=A4,leftMargin=57.6,rightMargin=57.6,topMargin=57.6,bottomMargin=57.6,title=doc["title"],author="").build(flow,onFirstPage=footer,onLaterPages=footer)


def _verify(stage, doc, outputs, links, deps):
    from openpyxl import load_workbook
    from pypdf import PdfReader
    checks = {"xlsx_count":0,"docx_count":0,"pdf_pages":0,"pdf_links":0,"source_count":len(doc["sources"]),"semantic_status":"unverified"}
    reader_text = []
    for path in outputs:
        if path.suffix == ".xlsx":
            wb = load_workbook(path)
            for ws in wb:
                for row in ws:
                    for cell in row:
                        if cell.data_type in ("f","e"):
                            fail("reporting_verification_failed", "Unexpected Excel formula/error")
                        if cell.value is not None:
                            reader_text.append(str(cell.value))
                        if isinstance(cell.value,float) and not math.isfinite(cell.value):
                            fail("reporting_verification_failed", "Nonfinite workbook value")
            checks["xlsx_count"] += 1
        elif path.suffix == ".docx":
            with zipfile.ZipFile(path) as z:
                root = ET.fromstring(z.read("word/document.xml"))
                reader_text.extend(e.text or "" for e in root.findall(".//w:t",NS))
                section = root.find(".//w:sectPr",NS)
                width = int(section.find("w:pgSz",NS).attrib["{"+NS["w"]+"}w"])
                margins = section.find("w:pgMar",NS).attrib
                usable = width-int(margins["{"+NS["w"]+"}left"])-int(margins["{"+NS["w"]+"}right"])
                for grid in root.findall(".//w:tblGrid",NS):
                    if sum(int(c.attrib["{"+NS["w"]+"}w"]) for c in grid) > usable+8:
                        fail("reporting_verification_failed", "Word table exceeds page width")
                bookmarks = {e.attrib.get("{"+NS["w"]+"}name") for e in root.findall(".//w:bookmarkStart", NS)}
                anchors = [e.attrib["{"+NS["w"]+"}anchor"] for e in root.findall(".//w:hyperlink", NS)
                           if "{"+NS["w"]+"}anchor" in e.attrib]
                if path.parent.name == "02_최종 보고서":
                    expected = [f"report_source_{next(s['reference_number'] for s in doc['references'] if s['id'] == sid)}"
                                for section in doc["sections"] for block in section["blocks"] for sid in block.get("source_ids", [])]
                    if anchors != expected or any(anchor not in bookmarks for anchor in anchors):
                        fail("reporting_verification_failed", "Word citation number/destination missing or mismatched")
                    markers = [(h.attrib["{"+NS["w"]+"}anchor"], "".join(e.text or "" for e in h.findall(".//w:t", NS)))
                               for h in root.findall(".//w:hyperlink", NS) if "{"+NS["w"]+"}anchor" in h.attrib]
                    if markers != [(anchor, "["+anchor.rsplit("_", 1)[1]+"]") for anchor in expected]:
                        fail("reporting_verification_failed", "Word citation label mismatch")
                    rels = ET.fromstring(z.read("word/_rels/document.xml.rels"))
                    by_rel = {r.attrib.get("Id"): r.attrib.get("Target") for r in rels}
                    targets = {by_rel.get(h.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
                               for h in root.findall(".//w:hyperlink", NS)}
                    for source in doc["references"]:
                        local = next(_source_targets([source["id"]], {source["id"]: source}, links))[1]
                        if local not in targets or (source.get("url") and source["url"] not in targets):
                            fail("reporting_verification_failed", "Word original URL/cell link missing")
            checks["docx_count"] += 1
    pdf = PdfReader(str(stage/"02_최종 보고서/보고서.pdf"))
    extracted = "\n".join(page.extract_text() or "" for page in pdf.pages)
    # PDF text extraction inserts line breaks when a legitimate title wraps.
    compact_title = re.sub(r"\s+", "", doc["title"])
    compact_extracted = re.sub(r"\s+", "", extracted)
    if not pdf.pages or not extracted.strip() or compact_title not in compact_extracted:
        fail("reporting_verification_failed", "PDF pages/text/title missing")
    reader_text.append(extracted)
    checks["pdf_pages"] = len(pdf.pages)
    uris = [str(a.get_object().get("/A",{}).get("/URI","")) for page in pdf.pages for a in page.get("/Annots",[])]
    checks["pdf_links"] = len([u for u in uris if u])
    page_refs = {page.indirect_reference.idnum for page in pdf.pages}
    internal = []
    for page in pdf.pages:
        for item in page.get("/Annots", []):
            annotation = item.get_object()
            action = annotation.get("/A", {})
            destination = annotation.get("/Dest") or (action.get("/D") if action.get("/S") == "/GoTo" else None)
            if destination is not None:
                if isinstance(destination, str):
                    destination = pdf.named_destinations.get(destination)
                    if destination is None:
                        fail("reporting_verification_failed", "PDF citation destination missing")
                    destination = destination.dest_array
                if not destination or getattr(destination[0], "idnum", None) not in page_refs:
                    fail("reporting_verification_failed", "PDF citation destination page missing")
                internal.append(destination)
    def outlines(items):
        for item in items:
            if isinstance(item, list):
                yield from outlines(item)
            else:
                yield item
    def destination_key(destination):
        return tuple((item.idnum, item.generation) if hasattr(item, "idnum")
                     else float(item) if isinstance(item, (int, float)) else str(item) for item in destination)
    entries = {entry.title: entry.dest_array for entry in outlines(pdf.outline)}
    reference_destinations = {}
    for source in doc["references"]:
        label = f"[{source['reference_number']}] "+source["title"]
        if label not in entries:
            fail("reporting_verification_failed", "PDF reference entry missing")
        reference_destinations[source["id"]] = destination_key(entries[label])
    expected = [reference_destinations[sid] for section in doc["sections"]
                for block in section["blocks"] for sid in block.get("source_ids", [])]
    if [destination_key(d) for d in internal] != expected:
        fail("reporting_verification_failed", "PDF citation number/destination missing or mismatched")
    checks["pdf_internal_links"] = len(internal)
    checks["reference_count"] = len(doc["references"])
    cited = {sid for section in doc["sections"] for block in section["blocks"] for sid in block.get("source_ids",[])}
    for sid in cited:
        expected = next(_source_targets([sid], {sid: next(s for s in doc["sources"] if s["id"] == sid)}, links))[1]
        if expected not in uris:
            fail("reporting_verification_failed", "PDF original-cell link missing")
    for source in doc["sources"]:
        link = links[source["source_id"]]
        wb = load_workbook(stage/"01_자료집"/link["file"])
        ws = wb[link["sheet"]]
        row = int(link["cell"][1:])
        rebuilt = "".join(ws.cell(row+i,5).value or "" for i in range(link["chunks"]))
        if rebuilt != source["text"]:
            fail("reporting_verification_failed", "Source text was truncated/changed")
    checks["structure_verified"] = True
    checks["visual_review"] = "not_run"
    checks["chart_editing"] = "rendered_images; data_cells_editable; regenerate_from_report_plan"
    return checks


def render_report(plan: dict, destination) -> dict:
    """Render to a new directory atomically; failures never advertise final files.

    Temporary charts and verification stay in a sibling scratch directory.
    Ready certifies file generation/structural checks, not semantic or visual QA.
    """
    stage = scratch = None
    target = None
    directory = str(destination)
    try:
        if not isinstance(destination, (str, Path)) or not str(destination).strip():
            fail("invalid_report_destination", "Destination must be a new local directory")
        target = Path(destination).absolute()
        if ".." in target.parts or str(target).startswith(("\\\\", "//")):
            fail("invalid_report_destination", "Destination cannot traverse or use a network path")
        target = target.resolve()
        directory = str(target)
        if target.exists():
            fail("report_destination_exists", "Destination already exists; existing results are preserved")
        normalized = validate_report_plan(plan)
        deps = _deps()
        doc = _adapt_plan(normalized)
        target.parent.mkdir(parents=True, exist_ok=True)
        stage = Path(tempfile.mkdtemp(prefix=".reader-stage-", dir=target.parent)).resolve()
        scratch = Path(tempfile.mkdtemp(prefix=".reader-scratch-", dir=target.parent)).resolve()
        library, final = stage/"01_자료집", stage/"02_최종 보고서"
        library.mkdir()
        final.mkdir()
        outputs, links = _write_libraries(doc, library, deps)
        by_id = {s["source_id"]:s for s in doc["sources"]}
        _write_docx(doc, final/"보고서.docx", by_id, links, scratch)
        _write_pdf(doc, final/"보고서.pdf", by_id, links)
        outputs.extend((final/"보고서.docx", final/"보고서.pdf"))
        checks = _verify(stage, doc, outputs, links, deps)
        descriptors = [{"path":str(target/path.relative_to(stage)), "kind":path.suffix[1:],
                        "title":path.stem} for path in outputs]
        if any(p.suffix not in (".xlsx", ".docx", ".pdf") for p in stage.rglob("*") if p.is_file()):
            fail("reporting_verification_failed", "Unexpected reader-folder file")
        stage.rename(target)
        stage = None
        return {"ready":True, "directory":directory, "files":descriptors, "checks":checks}
    except Exception as exc:
        code = exc.code if isinstance(exc, ToolError) else "reporting_generation_failed"
        return {"ready":False, "directory":directory, "files":[],
                "checks":{"structure_verified":False, "errors":[{"code":code,"message":str(exc)[:512]}],
                          "semantic_status":"unverified", "visual_review":"not_run"}}
    finally:
        for temporary in (stage, scratch):
            # Verify generated absolute paths before any recursive cleanup.
            if (temporary is not None and target is not None and
                    temporary.parent == target.parent and temporary.name.startswith(".reader-") and temporary.is_dir()):
                shutil.rmtree(temporary)
